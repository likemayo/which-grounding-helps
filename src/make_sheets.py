import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

import os
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
H = os.path.join(ROOT, "annotation-workbench.html")
h = open(H, encoding="utf-8").read()
i = h.index("const ALL = ") + len("const ALL = ")
items, _ = json.JSONDecoder().raw_decode(h[i:])

LABELS = ["MENTAL_TYPO", "KNOWLEDGE_GAP", "MISCONCEPTION",
          "STRUCTURAL_BLINDNESS", "SPEC_MISREADING", "INDETERMINATE"]

INK   = "1C1A19"
MUTED = "6B645E"
RULE  = "E2DDD7"
HEADBG = "2F4F43"
FILLIN = "FFFDE7"

def build(path, rows, who, scope_note):
    wb = Workbook()

    # ---------- 说明 ----------
    ws = wb.active
    ws.title = "说明"
    ws.column_dimensions["A"].width = 96
    lines = [
        ("LAK27 错误成因标注表", 15, True, INK),
        (f"标注者：{who}", 11, True, INK),
        ("", 10, False, INK),
        (f"范围：{scope_note}", 11, False, INK),
        ("配套工作台：https://claude.ai/code/artifact/ee4357c1-a7d7-4443-893a-bb07a69fea8a", 10, False, MUTED),
        ("", 10, False, INK),
        ("怎么填", 12, True, INK),
        ("1. 打开工作台，用筛选器切到你的范围，逐条看。", 11, False, INK),
        ("2. 在「标注」工作表里，只填 label 一列（浅黄底）。item_id 已按工作台顺序排好，不要重排或删行。", 11, False, INK),
        ("3. label 单元格是下拉框，六选一。拿不准就选 INDETERMINATE —— 它是真实类别，不是跳过键。", 11, False, INK),
        ("4. note 一列可选，只在你犹豫过的条目上写一句为什么，留给分歧会议用。", 11, False, INK),
        ("", 10, False, INK),
        ("两条硬规矩", 12, True, INK),
        ("· 标注期间不与另一位标注者讨论具体条目。讨论过再算的 kappa 不成立。", 11, False, INK),
        ("· 只填自己这份，不看对方的。", 11, False, INK),
        ("", 10, False, INK),
        ("六个类别", 12, True, INK),
        ("MENTAL_TYPO —— 想对了写岔了。代码其余部分证明他知道正确规则，缺陷只出现在部分适用处。", 11, False, INK),
        ("KNOWLEDGE_GAP —— 缺语法或 API，根本表达不出来。伸手去够一个不存在的东西。", 11, False, INK),
        ("MISCONCEPTION —— 持有一个具体的错误信念，代码忠实实现了它。必须能用一句「他以为……」说出来。", 11, False, INK),
        ("STRUCTURAL_BLINDNESS —— 每个表达式都站得住，错在排列。挪位置就能修好。", 11, False, INK),
        ("SPEC_MISREADING —— 干净地实现了另一条自洽规则，只是不是题目那条。你能说出那条规则是什么。", 11, False, INK),
        ("INDETERMINATE —— 证据不足以判定。", 11, False, INK),
        ("", 10, False, INK),
        ("编译器消息是症状，不是类别。", 12, True, "8C2F2A"),
        ("先读消息定位到哪一行出了什么，再回代码判断成因。cannot find symbol 可能是 KNOWLEDGE_GAP", 11, False, INK),
        ("（不知道有这个 API），也可能是 MENTAL_TYPO（变量名打错）—— 看这个符号在代码别处有没有被正确使用过。", 11, False, INK),
    ]
    for n, (txt, sz, bold, color) in enumerate(lines, start=1):
        c = ws.cell(row=n, column=1, value=txt)
        c.font = Font(name="Arial", size=sz, bold=bold, color=color)
        c.alignment = Alignment(wrap_text=True, vertical="top")

    r0 = len(lines) + 2
    ws.cell(row=r0, column=1, value="进度").font = Font(name="Arial", size=12, bold=True, color=INK)
    ws.cell(row=r0 + 1, column=1, value=f"已填 / 总数").font = Font(name="Arial", size=11, color=MUTED)
    pc = ws.cell(row=r0 + 1, column=2, value=f'=COUNTA(标注!B2:B{len(rows)+1})&" / {len(rows)}"')
    pc.font = Font(name="Arial", size=11, bold=True, color=INK)

    # ---------- 标注 ----------
    s = wb.create_sheet("标注")
    heads = ["item_id", "label", "note（可选）"]
    widths = [16, 26, 60]
    for j, (t, w) in enumerate(zip(heads, widths), start=1):
        c = s.cell(row=1, column=j, value=t)
        c.font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=HEADBG)
        c.alignment = Alignment(vertical="center")
        s.column_dimensions[get_column_letter(j)].width = w
    s.freeze_panes = "A2"

    thin = Side(style="thin", color=RULE)
    fill = PatternFill("solid", fgColor=FILLIN)
    for n, iid in enumerate(rows, start=2):
        a = s.cell(row=n, column=1, value=iid)
        a.font = Font(name="Arial", size=11, color=INK)
        for j in (2, 3):
            c = s.cell(row=n, column=j)
            c.font = Font(name="Arial", size=11, color=INK)
            c.fill = fill
        for j in (1, 2, 3):
            s.cell(row=n, column=j).border = Border(bottom=thin)

    dv = DataValidation(type="list", formula1='"' + ",".join(LABELS) + '"',
                        allow_blank=True, showDropDown=False)
    dv.error = "只能从六个类别里选一个"
    dv.errorTitle = "无效标签"
    s.add_data_validation(dv)
    dv.add(f"B2:B{len(rows)+1}")

    wb.save(path)
    print(path, len(rows), "条")


order = [x["id"] for x in items]
build(os.path.join(ROOT, "标注_Sophie_100.xlsx"),
      [x["id"] for x in items if x["s100"]],
      "Sophie Lin",
      "工作台筛选器选「双标注 100 条」，共 100 条。顺序与筛选后的顺序一致。")
build(os.path.join(ROOT, "标注_Yanlin_264.xlsx"),
      order,
      "Yanlin Wu",
      "工作台筛选器选「全部」，共 264 条（含 Sophie 也标的那 100 条）。顺序与工作台一致。")
